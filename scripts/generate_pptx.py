#!/usr/bin/env python3
"""
generate_pptx.py
MDファイルをPowerPoint(.pptx)に変換する。
# H1  → タイトルスライド
## H2 → セクション区切りスライド
### H3 + 本文 → コンテンツスライド
依存: python-pptx
インストール: pip install python-pptx
使用例:
  python scripts/generate_pptx.py \
    --input outputs/01_提案見積/01-2_提案書/技術提案書.md \
    --output outputs/01_提案見積/01-2_提案書/技術提案書.pptx
"""

import argparse
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("ERROR: python-pptx が必要です。pip install python-pptx を実行してください。")
    sys.exit(1)

# カラーパレット（GCインフラ案件向け）
COLOR_TITLE_BG = RGBColor(0x1F, 0x4E, 0x79)   # 濃紺
COLOR_SECTION_BG = RGBColor(0x2E, 0x75, 0xB6)  # 青
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_DARK = RGBColor(0x1F, 0x1F, 0x1F)
COLOR_ACCENT = RGBColor(0x00, 0x70, 0xC0)


def set_slide_background(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_slide(prs: Presentation, title: str, subtitle: str = ""):
    slide_layout = prs.slide_layouts[6]  # 空白レイアウト
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, COLOR_TITLE_BG)

    # タイトルテキストボックス
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(8.4), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER

    if subtitle:
        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(8.4), Inches(0.8))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = COLOR_WHITE
        p2.alignment = PP_ALIGN.CENTER
    return slide


def add_section_slide(prs: Presentation, title: str):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, COLOR_SECTION_BG)

    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(8.4), Inches(1.2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER
    return slide


def add_content_slide(prs: Presentation, title: str, content_lines: list[str]):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # タイトルバー
    titleBox = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(10), Inches(0.9))
    titleBox.fill.solid()
    titleBox.fill.fore_color.rgb = COLOR_TITLE_BG
    tf = titleBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.1)

    # コンテンツ
    contentBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(9.0), Inches(5.5))
    tf2 = contentBox.text_frame
    tf2.word_wrap = True

    for i, line in enumerate(content_lines):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        # 箇条書き（- や * で始まる行）
        if re.match(r"^[-*]\s+", line):
            p.text = re.sub(r"^[-*]\s+", "• ", line)
            p.font.size = Pt(16)
            p.level = 0
        elif re.match(r"^\s+[-*]\s+", line):
            p.text = re.sub(r"^\s+[-*]\s+", "  ◦ ", line)
            p.font.size = Pt(14)
            p.level = 1
        else:
            p.text = line
            p.font.size = Pt(16)
        p.font.color.rgb = COLOR_DARK
    return slide


def parse_md(md_text: str) -> list[dict]:
    """MDをスライド構造にパースする"""
    slides = []
    lines = md_text.splitlines()
    current = None

    for line in lines:
        h1 = re.match(r"^#\s+(.+)", line)
        h2 = re.match(r"^##\s+(.+)", line)
        h3 = re.match(r"^###\s+(.+)", line)

        if h1:
            if current:
                slides.append(current)
            current = {"type": "title", "title": h1.group(1), "content": []}
        elif h2:
            if current:
                slides.append(current)
            current = {"type": "section", "title": h2.group(1), "content": []}
        elif h3:
            if current:
                slides.append(current)
            current = {"type": "content", "title": h3.group(1), "content": []}
        else:
            if current and line.strip():
                current["content"].append(line)

    if current:
        slides.append(current)
    return slides


def main():
    parser = argparse.ArgumentParser(description="MD → PowerPoint変換")
    parser.add_argument("--input", "-i", required=True, help="入力MDファイルパス")
    parser.add_argument("--output", "-o", required=True, help="出力pptxファイルパス")
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        print(f"ERROR: 入力ファイルが見つかりません: {input_path}")
        sys.exit(1)

    md_text = input_path.read_text(encoding="utf-8")
    slides_data = parse_md(md_text)

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    for s in slides_data:
        if s["type"] == "title":
            subtitle = " ".join(s["content"][:1]) if s["content"] else ""
            add_title_slide(prs, s["title"], subtitle)
        elif s["type"] == "section":
            add_section_slide(prs, s["title"])
        else:
            add_content_slide(prs, s["title"], s["content"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    print(f"✅ 変換完了: {output_path}（スライド数: {len(prs.slides)}）")


if __name__ == "__main__":
    main()
